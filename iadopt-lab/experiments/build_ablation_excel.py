#!/usr/bin/env python
"""Render the shot-count ablation JSONL into its own Excel workbook.

Kept separate from the official reporting path on purpose: this reads only the ablation's
own results file and writes only into `experiments/output/`, so nothing it produces can be
mistaken for, or merged into, an official ranking.

Aggregation matches the official experiment — micro metrics via `aggregate_items`, summing
contributions across variables and computing F1 once — so the numbers mean the same thing
as the ones in the official rankings.
"""

from __future__ import annotations

import json
import statistics
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "experiments"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from shot_count_ablation import (  # noqa: E402
    EXPERIMENT,
    MODELS,
    OUT_DIR,
    RESULTS,
    SHOT_LEVELS,
    shot_pool,
)

from iadopt_eval.core import CLOSE_THRESHOLD, SCORER_VERSION, aggregate_items  # noqa: E402
from iadopt_lab.corpus.ingestion import load_canonical_records  # noqa: E402

WORKBOOK = OUT_DIR / "shot-count-ablation-results.xlsx"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F4F6F")


def _head(sheet, columns, widths=None):
    sheet.append(columns)
    for index, cell in enumerate(sheet[1], start=1):
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if widths:
            sheet.column_dimensions[cell.column_letter].width = widths[index - 1]
    sheet.freeze_panes = "A2"


def main() -> int:
    rows = [json.loads(line) for line in RESULTS.read_text(encoding="utf-8").splitlines() if line.strip()]
    grouped = defaultdict(list)
    for row in rows:
        grouped[(row["model_id"], row["shot_count"])].append(row)

    book = Workbook()

    # --- Sheet 1: summary, one row per model x shot count -------------------------
    summary = book.active
    summary.title = "Summary"
    _head(summary, [
        "Model", "Provider", "Reasoning", "Shots", "Evaluation N", "Scored", "Valid %",
        "Failed calls", "Close F1", "Close P", "Close R", "Exact F1", "Exact P", "Exact R",
        "Mean latency s", "Mean prompt tok", "Mean completion tok", "Total completion tok"],
        [22, 10, 14, 7, 13, 8, 9, 12, 10, 9, 9, 10, 9, 9, 14, 16, 19, 20])

    trend = defaultdict(dict)
    for model_id in MODELS:
        for shots in SHOT_LEVELS:
            batch = grouped.get((model_id, shots), [])
            if not batch:
                continue
            scored = [r for r in batch if r.get("evaluation")]
            ids = [r["variable_id"] for r in scored]
            agg = aggregate_items([r["evaluation"] for r in scored],
                                  expected_variable_ids=ids) if scored else None
            g = (lambda mode, metric: agg[mode]["metrics"][metric]["value"]) if agg else (lambda *a: 0.0)
            lat = [r["latency_seconds"] for r in batch if r.get("latency_seconds")]
            pin = [r["prompt_tokens"] for r in batch if r.get("prompt_tokens")]
            out = [r["completion_tokens"] for r in batch if r.get("completion_tokens")]
            trend[model_id][shots] = g("close", "f1")
            summary.append([
                model_id, "psnc", batch[0]["reasoning_mode"], shots, len(batch), len(scored),
                round(100 * len(scored) / len(batch), 1),
                sum(1 for r in batch if not r.get("ok")),
                round(g("close", "f1"), 4), round(g("close", "precision"), 4), round(g("close", "recall"), 4),
                round(g("exact", "f1"), 4), round(g("exact", "precision"), 4), round(g("exact", "recall"), 4),
                round(statistics.mean(lat), 2) if lat else None,
                round(statistics.mean(pin)) if pin else None,
                round(statistics.mean(out)) if out else None,
                sum(out) if out else None])

    # --- Sheet 2: per-variable ------------------------------------------------------
    detail = book.create_sheet("Per-variable")
    _head(detail, ["Model", "Shots", "Variable label", "Category", "Valid", "Close F1",
                   "Exact F1", "Latency s", "Completion tok", "Variable id"],
          [22, 7, 52, 20, 8, 10, 10, 11, 15, 68])
    for row in sorted(rows, key=lambda r: (r["model_id"], r["shot_count"], r.get("label") or "")):
        ev = row.get("evaluation")
        m = (lambda mode: ev[mode]["metrics"]["f1"]["value"]) if ev else (lambda mode: None)
        detail.append([row["model_id"], row["shot_count"], row.get("label"), row.get("category"),
                       bool(row.get("valid_prediction")),
                       round(m("close"), 4) if ev else None,
                       round(m("exact"), 4) if ev else None,
                       row.get("latency_seconds"), row.get("completion_tokens"),
                       row["variable_id"]])

    # --- Sheet 3: shot pool ---------------------------------------------------------
    poolsheet = book.create_sheet("Shot pool")
    _head(poolsheet, ["Position", "Used at shot levels", "Origin", "Variable label",
                      "Source path", "Variable id"], [10, 22, 30, 52, 56, 68])
    records = list(load_canonical_records(ROOT))
    for position, entry in enumerate(shot_pool(records), start=1):
        used = ", ".join(str(s) for s in SHOT_LEVELS if s >= position)
        origin = (f"official demonstration #{entry['demonstration_position']}"
                  if entry["demonstration_position"] else "added for ablation (sorted variable_id)")
        poolsheet.append([position, used, origin, entry.get("label"),
                          entry["source_path"], entry["variable_id"]])

    # --- Sheet 4: configuration -----------------------------------------------------
    conf = book.create_sheet("Configuration")
    _head(conf, ["Setting", "Qwen3.8-27B", "DeepSeek-V4-Flash"], [34, 52, 52])
    keys = ["prompt_variant", "temperature", "top_p", "max_output_tokens",
            "reasoning_mode", "reasoning_fields", "source"]
    labels = {"source": "Selected from (official result)"}
    for key in keys:
        conf.append([labels.get(key, key),
                     json.dumps(MODELS["Qwen3.8-27B"][key]) if key == "reasoning_fields" else MODELS["Qwen3.8-27B"][key],
                     json.dumps(MODELS["DeepSeek-V4-Flash"][key]) if key == "reasoning_fields" else MODELS["DeepSeek-V4-Flash"][key]])
    for label, value in [
        ("provider", "psnc"), ("experiment", EXPERIMENT),
        ("scorer_version", SCORER_VERSION), ("close_threshold", CLOSE_THRESHOLD),
        ("evaluation population", f"{len({r['variable_id'] for r in rows})} variables, identical at every shot level"),
        ("shot pool", "10 variables, nested prefixes, deterministic"),
        ("prompt fidelity", "byte-identical to the official renderer at 0/1/3/5 shots"),
        ("status", "EXPLORATORY side analysis - not part of the official experiment"),
    ]:
        conf.append([label, value, value])

    book.save(WORKBOOK)
    print(f"wrote {WORKBOOK.relative_to(ROOT)}")
    print()
    print("Close F1 by shot count:")
    for model_id, series in trend.items():
        line = "  ".join(f"{s}sh {series[s]:.3f}" for s in sorted(series))
        print(f"  {model_id:20} {line}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
